# -*- coding: utf-8 -*-
"""La hoja del día en Excel con las columnas que quiera cada centro (30-09-2026:
«poder eliminar hora llegada, hora carga, o tener las 3 o solo 2»)."""
import io

import openpyxl

import server

FILAS = [{"ruta": "CA_A1", "conductor": "ANA PEREZ", "movil": "612345678", "furgo": "1234ABC",
          "h_llegada": "08:40", "h_bajada": "09:00", "h_salida": "09:20", "observaciones": "ok"}]


def cabecera(xlsx):
    ws = openpyxl.load_workbook(io.BytesIO(xlsx)).active
    return [ws.cell(2, c).value for c in range(1, ws.max_column + 1)], ws


def generar(**k):
    return server._build_plantilla_excel(FILAS, set(), set(), "40", "30/09/2026", set(), set(), **k)


def test_sin_esconder_nada_sale_como_siempre():
    cab, ws = cabecera(generar())
    assert cab == ["RUTA", "CONDUCTOR", "MÓVIL", "VEHÍCULO", "H. LLEGADA", "H. CARGA", "H. SALIDA", "OBSERVACIONES"]
    assert [ws.cell(3, c).value for c in range(1, 9)] == [
        "CA_A1", "ANA PEREZ", "612345678", "1234ABC", "08:40", "09:00", "09:20", "ok"]
    assert "A1:D1" in [str(r) for r in ws.merged_cells.ranges]


def test_el_interruptor_viejo_sigue_quitando_llegada_y_carga():
    cab, ws = cabecera(generar(hide_nave_yard=True))
    assert cab == ["RUTA", "CONDUCTOR", "MÓVIL", "VEHÍCULO", "H. SALIDA", "OBSERVACIONES"]
    assert "A1:C1" in [str(r) for r in ws.merged_cells.ranges]


def test_cada_columna_por_su_cuenta():
    cab, ws = cabecera(generar(ocultar={"h_bajada"}))
    assert cab == ["RUTA", "CONDUCTOR", "MÓVIL", "VEHÍCULO", "H. LLEGADA", "H. SALIDA", "OBSERVACIONES"]
    assert [ws.cell(3, c).value for c in range(1, 8)] == [
        "CA_A1", "ANA PEREZ", "612345678", "1234ABC", "08:40", "09:20", "ok"]
    cab, _ = cabecera(generar(ocultar={"movil", "h_llegada"}))
    assert cab == ["RUTA", "CONDUCTOR", "VEHÍCULO", "H. CARGA", "H. SALIDA", "OBSERVACIONES"]


def test_el_movil_para_la_hoja():
    assert server._movil_para_hoja("+34612345678") == "612345678"
    assert server._movil_para_hoja("612 345 678") == "612345678"
    assert server._movil_para_hoja("+351912345678") == "+351912345678"
    assert server._movil_para_hoja("") == "" and server._movil_para_hoja(None) == ""


def test_no_se_pueden_quitar_las_columnas_de_siempre():
    cab, _ = cabecera(generar(ocultar={"ruta", "conductor", "furgo", "observaciones", "inventada"}))
    assert cab[:2] == ["RUTA", "CONDUCTOR"] and "VEHÍCULO" in cab and "OBSERVACIONES" in cab
