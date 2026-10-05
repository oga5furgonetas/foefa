# -*- coding: utf-8 -*-
"""Leer la hoja MENSUAL de Dani: fechas de verdad (con el 12 detrás de las
columnas de horas), colores del nombre, el rosa del día extra, la leyenda, las
máximas y las fórmulas sin valores guardados. Contra `cuadrante_excel.py`."""
import datetime as dt
import io

import openpyxl
import pytest
from openpyxl.styles import PatternFill

import cuadrante_excel as cx

AZUL, MELOCOTON, AMARILLO, ROSA = "FF00B0F0", "FFFABF8F", "FFFFFF00", "FFFF66CC"


def _relleno(c):
    return PatternFill("solid", fgColor=c)


def hoja(formulas=False, leyenda=True):
    """Una MENSUAL pequeña con la misma forma que la de septiembre-octubre."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MENSUAL"
    ws["B18"], ws["C18"] = "Orden", "Nombre"
    inicio = dt.date(2026, 9, 28)
    cols = list(range(4, 4 + 14))                 # D..Q: 28-09 .. 11-10
    for i, c in enumerate(cols):
        if formulas:
            ws.cell(18, c, "=DATE(2026,9,28)" if i == 0 else f"={openpyxl.utils.get_column_letter(c - 1)}18+1")
        else:
            ws.cell(18, c, dt.datetime.combine(inicio + dt.timedelta(days=i), dt.time()))
    ws.cell(18, 18, "Horas mes")                  # R: una columna de horas en medio
    ws.cell(18, 19, "TOTAL")                      # S
    if formulas:
        ws.cell(18, 20, "=Q18+1")                 # T: el 12-10, detrás de las de horas
    else:
        ws.cell(18, 20, dt.datetime(2026, 10, 12))
    filas = [
        ("1", "ANA PEREZ", None, ["1", "1", "N/T", "1", "1", "1", "N/T"] * 2 + ["1"]),
        ("2", "LUIS GOMEZ", AZUL, ["1"] * 5 + ["N/T", "N/T"] + ["1"] * 5 + ["N/T", "N/T", "1"]),
        ("3", "JOSE JORGE", AZUL, ["1", "1", "N/T", "1", "1", "1", "N/T"] * 2 + ["1"]),
        ("4", "SARA FINDE", MELOCOTON, ["N/T"] * 5 + ["1", "1"] + ["N/T"] * 5 + ["1", "1", "N/T"]),
        ("5", "PACO SEVA", AMARILLO, ["1"] * 4 + ["N/D"] * 11),
        ("6", "MARIA OFICINA", None, ["S"] * 15),
    ]
    for k, (orden, nombre, color, codigos) in enumerate(filas):
        r = 19 + k
        ws.cell(r, 2, int(orden))
        ws.cell(r, 3, nombre)
        if color:
            ws.cell(r, 3).fill = _relleno(color)
        for c, cod in zip(cols + [20], codigos):
            ws.cell(r, c, 1.0 if cod == "1" else cod)
    ws.cell(19, 5).fill = _relleno(ROSA)                  # Ana, 29-09: día extra
    ws.cell(19, 4, "V (05:43)")                           # Ana, 28-09: vacaciones con hora
    ws.cell(19, 20).fill = _relleno("FF00FF00")           # un 1 verde: el color NO cuenta
    ws.cell(30, 3, "RUTAS PLANIFICADAS (contador de 1)")
    ws.cell(31, 3, "MÁXIMAS AMAZON")
    for c in cols + [20]:
        ws.cell(31, c, 44)
    if leyenda:
        for r, (color, txt) in enumerate([(MELOCOTON, "solo trabajan sabado y domingo"),
                                          (AZUL, "gente de empresa que solo trabaja de lunes a viernes"),
                                          (AMARILLO, "Desde el dia 1 se van de inactividad"),
                                          (ROSA, "DÍA EXTRA (6º día de la semana)")], start=35):
            ws.cell(r, 27).fill = _relleno(color)
            ws.cell(r, 28, txt)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def por_nombre(d):
    return {f["nombre"]: f for f in d["filas"]}


def test_fechas_de_verdad_con_el_12_detras_de_las_columnas_de_horas():
    d = cx.leer_mensual(hoja())
    assert d["fechas"][0] == "2026-09-28" and d["fechas"][-1] == "2026-10-12"
    assert len(d["fechas"]) == 15 and not d["avisos"]


def test_sin_valores_guardados_calcula_las_formulas_de_la_cabecera():
    d = cx.leer_mensual(hoja(formulas=True))
    assert d["fechas"][0] == "2026-09-28" and d["fechas"][-1] == "2026-10-12"
    assert "2026-10-01" in d["fechas"]           # el 30 -> 1 del mes siguiente


def test_codigos_hora_y_extra():
    ana = por_nombre(cx.leer_mensual(hoja()))["ANA PEREZ"]["celdas"]
    assert ana["2026-09-28"] == {"cod": "V", "hora": "05:43", "extra": False}
    assert ana["2026-09-29"]["cod"] == "1" and ana["2026-09-29"]["extra"] is True
    assert ana["2026-10-12"]["extra"] is False    # un 1 de otro color no es extra


def test_tipos_por_el_color_del_nombre_y_patron_deducido():
    f = por_nombre(cx.leer_mensual(hoja()))
    assert f["LUIS GOMEZ"]["tipo"] == "empresa" and "patron" not in f["LUIS GOMEZ"]
    assert f["JOSE JORGE"]["tipo"] == "empresa" and f["JOSE JORGE"]["patron"] == [0, 1, 3, 4, 5]
    assert f["SARA FINDE"]["tipo"] == "finde"
    assert f["PACO SEVA"]["tipo"] == "normal" and f["PACO SEVA"]["inactivo_desde"] == "2026-10-01"
    assert f["ANA PEREZ"]["tipo"] == "normal" and f["ANA PEREZ"]["orden"] == 1


def test_la_oficina_no_entra_en_rutas():
    f = por_nombre(cx.leer_mensual(hoja()))
    assert f["MARIA OFICINA"]["tipo"] == "inactivo"


def test_maximas_y_la_fila_de_rutas_no_es_una_persona():
    d = cx.leer_mensual(hoja())
    assert d["maximas"]["2026-10-12"] == 44 and len(d["maximas"]) == 15
    assert "RUTAS PLANIFICADAS (contador de 1)" not in por_nombre(d)
    assert len(d["filas"]) == 6


def test_la_leyenda_de_la_hoja_manda():
    d = cx.leer_mensual(hoja())
    assert d["leyenda"] == {"FABF8F": "finde", "00B0F0": "empresa", "FFFF00": "sale", "FF66CC": "extra"}


def test_sin_leyenda_usa_los_colores_de_siempre():
    f = por_nombre(cx.leer_mensual(hoja(leyenda=False)))
    assert f["LUIS GOMEZ"]["tipo"] == "empresa" and f["SARA FINDE"]["tipo"] == "finde"


def _con_filas(filas):
    """Una MENSUAL mínima del 28-09 al 12-10 con las filas que se le den."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MENSUAL"
    ws["C18"] = "Nombre"
    for i in range(15):
        ws.cell(18, 4 + i, dt.datetime(2026, 9, 28) + dt.timedelta(days=i))
    for k, (nombre, codigos) in enumerate(filas):
        ws.cell(19 + k, 3, nombre)
        for i, cod in enumerate(codigos):
            if cod:
                ws.cell(19 + k, 4 + i, 1.0 if cod == "1" else cod)
    buf = io.BytesIO()
    wb.save(buf)
    return por_nombre(cx.leer_mensual(buf.getvalue()))


def test_quien_empieza_con_varios_n_t_seguidos_entra_el_primer_dia_de_trabajo():
    # Francisco Gil: N/T hasta el 05-10 y empieza el 06-10. Sin esto, el motor
    # lo ponía a trabajar del 1 al 4.
    f = _con_filas([("FRANCISCO GIL", ["N/T"] * 8 + ["1", "1", "1", "N/T", "N/T", "1", "1"])])
    assert f["FRANCISCO GIL"]["alta_desde"] == "2026-10-06"


def test_quien_empieza_con_su_dia_de_acompanante_entra_ese_dia():
    f = _con_filas([("CAROLINA NUEVA", ["", "N/T", "N/T", "RA", "1", "N/T", "1"] + [""] * 8)])
    assert f["CAROLINA NUEVA"]["alta_desde"] == "2026-10-01"


def test_quien_ya_trabajaba_no_tiene_fecha_de_entrada():
    f = _con_filas([("ANA DE SIEMPRE", ["1", "1", "N/T", "N/T", "1", "1", "1"] * 2 + ["1"]),
                    ("LIBRA UNO Y EMPIEZA", ["N/T", "1", "1", "1", "1", "1", "N/T"] + [""] * 8)])
    assert "alta_desde" not in f["ANA DE SIEMPRE"]
    assert "alta_desde" not in f["LIBRA UNO Y EMPIEZA"]


def test_una_hoja_que_no_es_el_cuadrante_da_un_error_claro():
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Hola"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="Nombre"):
        cx.leer_mensual(buf.getvalue())


def _semanal_de_otro_centro():
    """Como la de DGA2 (05-10-2026): la primera pestaña no es el cuadrante y la semanal
    trae solo el NÚMERO del día (7..18 y 29, 30, 1, 2…), sin mes."""
    wb = openpyxl.Workbook()
    d = wb.active
    d.title = "diario"
    d["A1"] = "4"
    d["E10"], d["F10"] = "Nombre", "Código"
    ws = wb.create_sheet("semanal")
    ws["A5"], ws["B5"] = "Orden", "Nombre"
    dias = list(range(28, 31)) + list(range(1, 9))          # 28, 29, 30, 1..8 (cambia de mes)
    for i, n in enumerate(dias):
        ws.cell(5, 3 + i, float(n))
        ws.cell(6, 3 + i, "1" if i % 2 == 0 else " N/T")
    ws["A6"], ws["B6"] = 1, "ADRIAN SALAZAR"
    ws["B7"] = "MAXIMA"
    return wb


def _bytes(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_hoja_semanal_con_solo_numero_de_dia_pone_el_mes_y_avisa():
    r = cx.leer_mensual(_bytes(_semanal_de_otro_centro()), ref=dt.date(2026, 10, 5))
    assert r["fechas"][0] == "2026-09-28" and r["fechas"][-1] == "2026-10-08"
    assert len(r["fechas"]) == 11
    assert any("solo trae el número del día" in a and "28/09/2026 al 08/10/2026" in a for a in r["avisos"])
    assert r["filas"][0]["celdas"]["2026-10-01"]["cod"] in ("1", "N/T")
    assert r["filas"][0]["celdas"]["2026-09-28"]["cod"] == "1"       # el de la columna 28


def test_se_prueban_todas_las_pestanas_aunque_la_primera_no_sirva():
    r = cx.leer_mensual(_bytes(_semanal_de_otro_centro()), ref=dt.date(2026, 10, 5))
    assert r["filas"] and r["filas"][0]["nombre"] == "ADRIAN SALAZAR"


def test_si_ninguna_pestana_tiene_dias_da_el_error_de_siempre():
    wb = openpyxl.Workbook()
    wb.active["B2"], wb.active["C2"] = "Orden", "Nombre"
    with pytest.raises(ValueError, match="no tiene fechas"):
        cx.leer_mensual(_bytes(wb))


def test_entre_varias_semanales_gana_la_que_llega_mas_tarde():
    wb = openpyxl.Workbook()
    for nombre, dias in (("semanal", [20, 21, 22, 23, 24, 25, 26]), ("semanal 2", [29, 30, 1, 2, 3, 4, 5])):
        ws = wb.active if nombre == "semanal" else wb.create_sheet(nombre)
        ws.title = nombre
        ws["A5"], ws["B5"] = "Orden", "Nombre"
        for k, n in enumerate(dias):
            ws.cell(5, 3 + k, float(n))
        ws["B6"] = "ADRIAN SALAZAR"
    r = cx.leer_mensual(_bytes(wb), ref=dt.date(2026, 10, 5))           # la más reciente: hasta el 05-10
    assert r["fechas"][0] == "2026-09-29" and r["fechas"][-1] == "2026-10-05"
