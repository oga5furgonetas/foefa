# -*- coding: utf-8 -*-
"""La hoja del cuadrante (30-09-2026, Dani: «que si le clicas se expanda y sea
un propio Excel con su diario, su mensual y lo de deber horas, y se pueda
editar»): el rosa escrito a mano («1E»), Ctrl+Z que solo deshace lo suyo, los
apuntes del DIARIO y el registro de DEBER HORAS, con su importación."""
import datetime as dt
import io
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import HTTPException

import cuadrante_excel as cx
import server
from base_de_mentira import Col, Db

DANI = {"sub": "a1", "role": "admin", "name": "Dani", "sa": True}


@pytest.fixture
def base(monkeypatch):
    db = Db()
    db["drivers"] = Col([
        {"id": "d1", "name": "ANA PÉREZ", "center": "OGA5", "active": True},
        {"id": "d2", "name": "KARIM ERRIFAI HADDAOUI", "center": "OGA5", "active": True},
    ])
    auditoria = []

    async def audit(user, accion, detalle=None):
        auditoria.append(accion)

    monkeypatch.setattr(server, "db", db)
    monkeypatch.setattr(server, "_audit", audit)
    return SimpleNamespace(db=db, auditoria=auditoria)


def _turno(db, did, fecha):
    return next((d for d in db["shifts"].docs if d["driver_id"] == did and d["date"] == fecha), None)


# ── MENSUAL: EL ROSA A MANO Y CTRL+Z ─────────────────────────────────────────

@pytest.mark.asyncio(loop_scope="session")
async def test_1e_es_un_dia_de_trabajo_en_rosa_y_en_un_libre_no_hay_rosa(base):
    r = await server.cuadrante_celdas({"center": "OGA5", "cambios": [
        {"driver_id": "d1", "fecha": "2026-10-05", "cod": "1", "extra": True},
        {"driver_id": "d2", "fecha": "2026-10-05", "cod": "N/T", "extra": True}]}, DANI)
    assert r["guardadas"] == 2 and r["cambio_id"]
    assert _turno(base.db, "d1", "2026-10-05")["extra"] is True
    assert _turno(base.db, "d2", "2026-10-05")["extra"] is False


@pytest.mark.asyncio(loop_scope="session")
async def test_ctrl_z_solo_deshace_lo_ultimo_si_es_lo_suyo(base):
    r1 = await server.cuadrante_celdas({"center": "OGA5", "cambios": [
        {"driver_id": "d1", "fecha": "2026-10-06", "cod": "1"}]}, DANI)
    # Entre medias, otra persona cambia otra celda.
    base.db["cuadrante_cambios"].docs[-1]["en"] = "2026-09-30T10:00:00"
    r2 = await server.cuadrante_celdas({"center": "OGA5", "cambios": [
        {"driver_id": "d2", "fecha": "2026-10-06", "cod": "V"}]}, {**DANI, "name": "Luis"})
    base.db["cuadrante_cambios"].docs[-1]["en"] = "2026-09-30T10:05:00"
    with pytest.raises(HTTPException) as e:
        await server.cuadrante_deshacer({"center": "OGA5", "solo_si": r1["cambio_id"]}, DANI)
    assert e.value.status_code == 409 and "Luis" in e.value.detail
    assert _turno(base.db, "d1", "2026-10-06")["cod"] == "1"            # no se ha tocado nada
    await server.cuadrante_deshacer({"center": "OGA5", "solo_si": r2["cambio_id"]}, DANI)
    assert _turno(base.db, "d2", "2026-10-06") is None                   # lo último, deshecho


# ── DIARIO ───────────────────────────────────────────────────────────────────

@pytest.mark.asyncio(loop_scope="session")
async def test_los_apuntes_del_dia_se_guardan_y_se_borran(base):
    await server.cuadrante_diario_guardar({"center": "OGA5", "fecha": "2026-09-30", "notas": "Illa: 4455NKC", "cambios": [
        {"driver_id": "d1", "campo": "nota", "valor": "2wave"},
        {"driver_id": "d1", "campo": "roster", "valor": "L1 - AZUL - 3 DIAS"},
        {"driver_id": "d2", "campo": "libres", "valor": "ALVARO FLORES 18/10"},
        {"driver_id": "x.y", "campo": "nota", "valor": "no"},              # un id con punto no entra
        {"driver_id": "d2", "campo": "otra", "valor": "no"}]}, DANI)       # ni un campo que no existe
    v = await server.cuadrante_diario_ver("OGA5", "2026-09-30", DANI)
    assert v["notas"] == "Illa: 4455NKC"
    assert v["filas"] == {"d1": {"nota": "2wave", "roster": "L1 - AZUL - 3 DIAS"}, "d2": {"libres": "ALVARO FLORES 18/10"}}
    await server.cuadrante_diario_guardar({"center": "OGA5", "fecha": "2026-09-30", "cambios": [
        {"driver_id": "d1", "campo": "nota", "valor": ""}]}, DANI)
    v = await server.cuadrante_diario_ver("OGA5", "2026-09-30", DANI)
    assert v["filas"]["d1"] == {"roster": "L1 - AZUL - 3 DIAS"} and v["notas"] == "Illa: 4455NKC"
    assert (await server.cuadrante_diario_ver("OGA5", "2026-10-01", DANI))["filas"] == {}


# ── DEBER HORAS ──────────────────────────────────────────────────────────────

@pytest.mark.asyncio(loop_scope="session")
async def test_una_fila_se_liga_a_su_ficha_solo_con_el_nombre_entero(base):
    a = await server.deber_horas_anadir({"center": "OGA5", "fecha": "2026-08-30", "persona": "karim errifai haddaoui",
                                         "cuanto": "1 día"}, DANI)
    b = await server.deber_horas_anadir({"center": "OGA5", "fecha": "2026-08-30", "persona": "brao"}, DANI)
    assert a["driver_id"] == "d2" and b["driver_id"] is None               # «brao» se queda como texto
    c = await server.deber_horas_editar(b["id"], {"persona": "ANA PÉREZ", "saldado": "reten 12/10"}, DANI)
    assert c["driver_id"] == "d1" and c["saldado"] == "reten 12/10"
    filas = (await server.deber_horas_ver("OGA5", DANI))["filas"]
    assert [f["persona"] for f in filas] == ["karim errifai haddaoui", "ANA PÉREZ"]
    await server.deber_horas_borrar(a["id"], DANI)
    assert len((await server.deber_horas_ver("OGA5", DANI))["filas"]) == 1
    assert "deber_horas_borrada" in base.auditoria
    with pytest.raises(HTTPException):
        await server.deber_horas_anadir({"center": "OGA5", "persona": "  "}, DANI)


def _libro():
    wb = openpyxl.Workbook()
    wb.active.title = "MENSUAL"
    ws = wb.create_sheet("deber horas")
    ws["B3"] = "DIAS"
    ws["B5"], ws["D5"] = "PABLO OTERO", "fichado reten 12/07"
    ws["B8"] = dt.datetime(2026, 7, 7)
    ws["B9"], ws["D9"] = "LORENZO MOREY MESQUIDA", "reten 28/07"
    ws["B10"], ws["C10"] = "ZEUS PEÑA", "2 dias 27-28/07"
    ws["B12"] = dt.datetime(2026, 8, 30)
    ws["B13"] = "KARIM ERRIFAI HADDAOUI"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_lee_la_hoja_deber_horas_como_esta():
    h = cx.leer_deber_horas(_libro())
    assert h["hoja"] == "deber horas"
    assert [(f["fecha"], f["persona"], f["nota"], f["saldado"]) for f in h["filas"]] == [
        ("", "PABLO OTERO", "", "fichado reten 12/07"),
        ("2026-07-07", "LORENZO MOREY MESQUIDA", "", "reten 28/07"),
        ("2026-07-07", "ZEUS PEÑA", "2 dias 27-28/07", ""),
        ("2026-08-30", "KARIM ERRIFAI HADDAOUI", "", "")]


def test_un_excel_sin_esa_hoja_da_un_error_claro():
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError) as e:
        cx.leer_deber_horas(buf.getvalue())
    assert "deber horas" in str(e.value)


class _Fichero:
    def __init__(self, datos):
        self.datos = datos

    async def read(self):
        return self.datos


@pytest.mark.asyncio(loop_scope="session")
async def test_importar_dos_veces_no_duplica(base):
    previa = await server.deber_horas_importar(_Fichero(_libro()), "OGA5", "", DANI)
    assert previa["leidas"] == 4 and previa["nuevas"] == 4 and previa["pendientes"] == 2 and not previa["guardado"]
    assert base.db["cuadrante_deber_horas"].docs == []                    # la previa no guarda nada
    r = await server.deber_horas_importar(_Fichero(_libro()), "OGA5", "1", DANI)
    assert r["guardado"] and len(base.db["cuadrante_deber_horas"].docs) == 4
    karim = next(x for x in base.db["cuadrante_deber_horas"].docs if x["persona"].startswith("KARIM"))
    assert karim["driver_id"] == "d2" and karim["origen"] == "excel"
    otra = await server.deber_horas_importar(_Fichero(_libro()), "OGA5", "1", DANI)
    assert otra["nuevas"] == 0 and len(base.db["cuadrante_deber_horas"].docs) == 4
